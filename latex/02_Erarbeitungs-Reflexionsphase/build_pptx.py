# -*- coding: utf-8 -*-
"""Generates the Phase 2 presentation as a native, editable PowerPoint file,
mirroring the content and visual style of presentation.tex (metropolis theme:
dark header bar, teal accent, rounded screenshot cards)."""

import re
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from PIL import Image
import os

BASE = r"c:\Users\tizia\Documents\GitHub\IUDatabaseProject\latex\02_Erarbeitungs-Reflexionsphase"
SHOTS = os.path.join(BASE, "screenshots")
OUT = os.path.join(BASE, "Senger-Tizian_IU14143428_DLBDSPBDM01_D_Erarbeitungs-Reflexionsphase_PR.pptx")

DARK = RGBColor(0x1F, 0x2A, 0x2E)
TEAL = RGBColor(0x00, 0x79, 0x6B)
LIGHT_BG = RGBColor(0xFA, 0xFA, 0xFA)
CODE_BG = RGBColor(0xF2, 0xF2, 0xF0)
CODE_BORDER = RGBColor(0xD0, 0xD0, 0xCE)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREY_TEXT = RGBColor(0x33, 0x33, 0x33)
ORANGE_STR = RGBColor(0xB0, 0x5A, 0x00)

SQL_KEYWORDS = set("""CREATE TABLE TRIGGER PRIMARY KEY AUTOINCREMENT NOT NULL UNIQUE DEFAULT
CHECK REFERENCES ON DELETE CASCADE VARCHAR INTEGER CHAR DATE TIME DATETIME DECIMAL
BOOLEAN IN BETWEEN AND OR SELECT FROM WHERE JOIN GROUP BY ORDER DESC ASC AS
BEFORE INSERT FOR EACH ROW WHEN BEGIN RAISE ABORT END VALUES UPDATE SET IS""".split())

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

SW, SH = prs.slide_width, prs.slide_height


def new_slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    bg.fill.solid()
    bg.fill.fore_color.rgb = LIGHT_BG
    bg.line.fill.background()
    bg.shadow.inherit = False
    s.shapes._spTree.remove(bg._element)
    s.shapes._spTree.insert(2, bg._element)
    return s


def add_header(slide, title, page_no=None, total=None):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, Inches(0.95))
    bar.fill.solid()
    bar.fill.fore_color.rgb = DARK
    bar.line.fill.background()
    bar.shadow.inherit = False
    tb = bar.text_frame
    tb.margin_left = Inches(0.4)
    tb.margin_top = Inches(0.08)
    tb.word_wrap = True
    tb.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tb.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.size = Pt(26)
    r.font.bold = True
    r.font.color.rgb = WHITE
    r.font.name = "Calibri"
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(0.95), Inches(2.6), Pt(3))
    accent.fill.solid()
    accent.fill.fore_color.rgb = TEAL
    accent.line.fill.background()
    accent.shadow.inherit = False
    if page_no is not None:
        pb = slide.shapes.add_textbox(SW - Inches(1.3), SH - Inches(0.45), Inches(1.1), Inches(0.35))
        pf = pb.text_frame
        pf.paragraphs[0].alignment = PP_ALIGN.RIGHT
        r2 = pf.paragraphs[0].add_run()
        r2.text = f"{page_no}/{total}"
        r2.font.size = Pt(11)
        r2.font.color.rgb = RGBColor(0x88, 0x88, 0x88)


def add_code_box(slide, left, top, width, height, code_text, font_size=11):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    box.adjustments[0] = 0.04
    box.fill.solid()
    box.fill.fore_color.rgb = CODE_BG
    box.line.color.rgb = CODE_BORDER
    box.line.width = Pt(0.75)
    box.shadow.inherit = False
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.12)
    tf.margin_right = Inches(0.1)
    tf.margin_top = Inches(0.08)
    tf.margin_bottom = Inches(0.08)
    lines = code_text.strip("\n").split("\n")
    first = True
    for line in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(0)
        p.space_before = Pt(0)
        tokens = re.split(r"(\s+|[(),;])", line)
        for tok in tokens:
            if tok == "":
                continue
            run = p.add_run()
            run.text = tok
            run.font.name = "Consolas"
            run.font.size = Pt(font_size)
            up = tok.upper().strip("(),;")
            if up in SQL_KEYWORDS and tok.strip():
                run.font.color.rgb = TEAL
                run.font.bold = True
            elif re.match(r"^'.*'$", tok.strip()):
                run.font.color.rgb = ORANGE_STR
            elif tok.strip().startswith("--"):
                run.font.color.rgb = RGBColor(0x77, 0x77, 0x77)
                run.font.italic = True
            else:
                run.font.color.rgb = GREY_TEXT
    return box


def add_text(slide, left, top, width, height, runs, font_size=14, bold_first=False, italic=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    if isinstance(runs, str):
        runs = [(runs, {})]
    for i, (text, style) in enumerate(runs):
        r = p.add_run()
        r.text = text
        r.font.size = Pt(style.get("size", font_size))
        r.font.bold = style.get("bold", bold_first and i == 0)
        r.font.italic = style.get("italic", italic)
        r.font.name = style.get("font", "Calibri")
        r.font.color.rgb = style.get("color", GREY_TEXT)
    return box


def add_bullets(slide, left, top, width, height, items, font_size=16, sub_font_size=13):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for item in items:
        if isinstance(item, tuple):
            text, level = item
        else:
            text, level = item, 0
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = level
        p.space_after = Pt(6)
        bullet = "•  " if level == 0 else "-  "
        r = p.add_run()
        r.text = bullet + text
        r.font.size = Pt(font_size if level == 0 else sub_font_size)
        r.font.color.rgb = GREY_TEXT
        r.font.name = "Calibri"
    return box


def add_screenshot(slide, path, left, top, max_w, max_h, caption=None):
    frame = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, max_w, max_h)
    frame.adjustments[0] = 0.06
    frame.fill.solid()
    frame.fill.fore_color.rgb = WHITE
    frame.line.color.rgb = CODE_BORDER
    frame.line.width = Pt(0.75)
    frame.shadow.inherit = False
    frame.text_frame.paragraphs[0].text = ""
    if path and os.path.exists(path):
        im = Image.open(path)
        iw, ih = im.size
        ratio = iw / ih
        pad = Inches(0.12)
        avail_w, avail_h = max_w - 2 * pad, max_h - 2 * pad
        if avail_w / avail_h > ratio:
            disp_h = avail_h
            disp_w = int(disp_h * ratio)
        else:
            disp_w = avail_w
            disp_h = int(disp_w / ratio)
        pic_left = left + (max_w - disp_w) // 2
        pic_top = top + (max_h - disp_h) // 2
        slide.shapes.add_picture(path, pic_left, pic_top, width=disp_w, height=disp_h)
    else:
        tf = frame.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = "Screenshot fehlt: " + (caption or "")
        r.font.size = Pt(11)
        r.font.italic = True
        r.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
    return frame


TOTAL = 25  # actual slide count of this deck (pptx has no separate section-divider slides)

# ---------------------------------------------------------------------------
# Folie 1: Titel
# ---------------------------------------------------------------------------
s = new_slide()
add_text(s, Inches(0.8), Inches(2.3), Inches(11.7), Inches(1.0),
         [("Datenbankentwurf für eine Buchtausch-App", {})],
         font_size=34, bold_first=True, )
add_text(s, Inches(0.8), Inches(3.1), Inches(11.7), Inches(0.6),
         [("Portfolioteil 2 -- Erarbeitungs-/Reflexionsphase", {"color": TEAL, "size": 20})])
line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(3.75), Inches(6), Pt(1.5))
line.fill.solid(); line.fill.fore_color.rgb = RGBColor(0xCC, 0xCC, 0xCC); line.line.fill.background(); line.shadow.inherit = False
add_text(s, Inches(0.8), Inches(4.1), Inches(11.7), Inches(1.6),
         [("Tizian Senger\n", {"size": 16}),
          ("DLBDSPBDM01_D -- Data-Mart-Erstellung in SQL\n", {"size": 13, "color": RGBColor(0x66,0x66,0x66)}),
          ("Matrikelnummer: IU14143428\n", {"size": 13, "color": RGBColor(0x66,0x66,0x66)}),
          ("https://github.com/TizianSenger/IUDatabaseProject", {"size": 13, "color": RGBColor(0x66,0x66,0x66)}),
          ])

# ---------------------------------------------------------------------------
# Folie 2: Agenda
# ---------------------------------------------------------------------------
s = new_slide()
add_header(s, "Agenda", 2, TOTAL)
add_bullets(s, Inches(1.0), Inches(1.6), Inches(10), Inches(4.5), [
    "Überblick",
    "Tabellen und SQL-Statements",
    "Dreifachbeziehungen",
    "Constraints und Geschäftsregeln",
    "Zusammenfassung und Reflexion",
], font_size=20)

# ---------------------------------------------------------------------------
# Folie 3: Überblick
# ---------------------------------------------------------------------------
s = new_slide()
add_header(s, "Überblick über das Datenbankdesign", 3, TOTAL)
add_bullets(s, Inches(0.7), Inches(1.25), Inches(11.9), Inches(5.5), [
    ("Verwendetes DBMS: SQLite 3 -- serverlos, eine einzelne Datei, voll SQL-fähig "
     "(CHECK, FOREIGN KEY, TRIGGER, JOIN); laut Aufgabenstellung frei wählbar.", 0),
    ("Anzahl Tabellen: 17 (13 Entitäten + 2 Dreifachbeziehungstabellen ausleihe/verfuegbarkeit "
     "+ 2 M:N-Tabellen buch_autor/buch_genre).", 0),
    ("Normalform: 3. Normalform (3NF) -- keine transitiven Abhängigkeiten, "
     "Autor/Verlag/Genre/Sprache konsequent ausgelagert.", 0),
    ("Gegenüber Phase 1 umgesetztes Tutor-Feedback:", 0),
    ("BUCH (Werk) und BUCHEXEMPLAR (physische Kopie inkl. Eigentümer:in) getrennt.", 1),
    ("BEWERTUNG per eindeutigem FK an ausleihe_id gekoppelt statt eigener Dreifachbeziehung.", 1),
    ("Durchgängige CHECK-Constraints (Sterne 1-5, Versandkosten >= 0, startzeit < endzeit).", 1),
    ("Negativtestfälle für alle zentralen Geschäftsregeln (siehe Folien 21-22).", 1),
], font_size=16, sub_font_size=14)


def add_entity_slide(title, page_no, create_sql, testfall_runs, shot_file, shot_caption, section_label="Tabelle"):
    s = new_slide()
    add_header(s, title, page_no, TOTAL)
    left_x, left_w = Inches(0.5), Inches(6.15)
    right_x, right_w = Inches(6.9), Inches(5.95)
    top = Inches(1.2)
    add_text(s, left_x, top, left_w, Inches(0.35), [("CREATE-Statement:", {"bold": True, "size": 15})])
    add_code_box(s, left_x, top + Inches(0.4), left_w, Inches(4.5), create_sql, font_size=11.5)

    add_text(s, right_x, top, right_w, Inches(1.5), testfall_runs, font_size=14)
    add_screenshot(s, os.path.join(SHOTS, shot_file) if shot_file else None,
                    right_x, top + Inches(1.65), right_w, Inches(3.3), shot_caption)
    return s


entities = [
    dict(title="Tabelle: NUTZER", page=4, shot="nutzer.png",
         caption="Tabellenergebnis mit 11 Nutzenden der Rolle 'nutzer'",
         sql="""CREATE TABLE nutzer (
    nutzer_id      INTEGER PRIMARY KEY AUTOINCREMENT,
    vorname        VARCHAR(50)  NOT NULL,
    nachname       VARCHAR(50)  NOT NULL,
    email          VARCHAR(100) NOT NULL UNIQUE,
    passwort_hash  VARCHAR(255) NOT NULL,
    telefon        VARCHAR(20),
    rolle          VARCHAR(10)  NOT NULL DEFAULT 'nutzer'
        CHECK (rolle IN ('nutzer', 'admin')),
    registriert_am DATE NOT NULL DEFAULT CURRENT_DATE
);""",
         testfall=[("Testfall: ", {"bold": True}),
                    ("SELECT nutzer_id, vorname, nachname, email FROM nutzer WHERE rolle = 'nutzer';", {"font": "Consolas", "size": 11}),
                    ("  (11 Treffer)", {})]),
    dict(title="Tabelle: ADRESSE", page=5, shot="adresse.png",
         caption="Ergebniszeile mit Anna Bauers Adresse",
         sql="""CREATE TABLE adresse (
    adresse_id  INTEGER PRIMARY KEY AUTOINCREMENT,
    nutzer_id   INTEGER NOT NULL REFERENCES nutzer(nutzer_id) ON DELETE CASCADE,
    strasse     VARCHAR(100) NOT NULL,
    hausnummer  VARCHAR(10)  NOT NULL,
    plz         VARCHAR(10)  NOT NULL,
    stadt       VARCHAR(50)  NOT NULL,
    land        VARCHAR(50)  NOT NULL DEFAULT 'Deutschland'
);""",
         testfall=[("Testfall: ", {"bold": True}),
                    ("Adresse von Anna Bauer per JOIN adresse--nutzer → Musterstrasse 12, 10115 Berlin.", {})]),
    dict(title="Tabelle: VERLAG", page=6, shot="verlag.png",
         caption="Liste der 9 deutschen Verlage",
         sql="""CREATE TABLE verlag (
    verlag_id  INTEGER PRIMARY KEY AUTOINCREMENT,
    name       VARCHAR(100) NOT NULL UNIQUE,
    land       VARCHAR(50)
);""",
         testfall=[("Testfall: ", {"bold": True}),
                    ("SELECT name FROM verlag WHERE land = 'Deutschland';", {"font": "Consolas", "size": 11}),
                    ("  (9 von 10 Verlagen, Diogenes ist Schweizer Verlag)", {})]),
    dict(title="Tabelle: SPRACHE", page=7, shot="sprache.png",
         caption="Ergebniszeile ('Deutsch', 'DE')",
         sql="""CREATE TABLE sprache (
    sprache_id   INTEGER PRIMARY KEY AUTOINCREMENT,
    bezeichnung  VARCHAR(50) NOT NULL UNIQUE,
    iso_code     CHAR(2) NOT NULL UNIQUE
);""",
         testfall=[("Testfall: ", {"bold": True}),
                    ("SELECT bezeichnung, iso_code FROM sprache WHERE bezeichnung = 'Deutsch';", {"font": "Consolas", "size": 11}),
                    ("  → ('Deutsch', 'DE')", {})]),
    dict(title="Tabelle: GENRE", page=8, shot="genre.png",
         caption="Liste der 5 Roman-Buecher",
         sql="""CREATE TABLE genre (
    genre_id     INTEGER PRIMARY KEY AUTOINCREMENT,
    bezeichnung  VARCHAR(50) NOT NULL UNIQUE
);""",
         testfall=[("Testfall: ", {"bold": True}),
                    ("Alle Bücher im Genre „Roman“ per JOIN buch_genre -- 5 Treffer (Die Verwandlung, Der Steppenwolf, Der Zauberberg, Das Geisterhaus, Schachnovelle).", {})]),
    dict(title="Tabelle: AUTOR", page=9, shot="autor.png",
         caption="Ergebnis: Die Verwandlung",
         sql="""CREATE TABLE autor (
    autor_id  INTEGER PRIMARY KEY AUTOINCREMENT,
    vorname   VARCHAR(50) NOT NULL,
    nachname  VARCHAR(50) NOT NULL
);""",
         testfall=[("Testfall: ", {"bold": True}),
                    ("Alle Bücher von Franz Kafka per JOIN buch_autor → „Die Verwandlung“.", {})]),
    dict(title="Tabelle: BUCH", page=10, shot="buch.png",
         caption="Tabelle mit 12 Buechern inkl. Verlag und Sprache",
         sql="""CREATE TABLE buch (
    buch_id          INTEGER PRIMARY KEY AUTOINCREMENT,
    titel            VARCHAR(150) NOT NULL,
    verlag_id        INTEGER REFERENCES verlag(verlag_id),
    sprache_id       INTEGER NOT NULL REFERENCES sprache(sprache_id),
    erscheinungsjahr INTEGER NOT NULL
        CHECK (erscheinungsjahr BETWEEN 1450 AND 2100)
);""",
         testfall=[("Testfall: ", {"bold": True}),
                    ("Werk-Stammdaten inkl. Verlag/Sprache per JOIN -- 12 Bücher, z. B. ('Die Verwandlung', 'Fischer Verlag', 'Deutsch', 1915).", {})]),
    dict(title="Tabelle: BUCHEXEMPLAR", page=11, shot="buchexemplar.png",
         caption="Zwei Exemplare von 1984 mit unterschiedlichen Eigentuemer:innen",
         sql="""CREATE TABLE buchexemplar (
    exemplar_id     INTEGER PRIMARY KEY AUTOINCREMENT,
    buch_id         INTEGER NOT NULL REFERENCES buch(buch_id),
    eigentuemer_id  INTEGER NOT NULL REFERENCES nutzer(nutzer_id) ON DELETE CASCADE,
    zustand         VARCHAR(20) NOT NULL
        CHECK (zustand IN ('neu','gut','gebraucht','stark gebraucht')),
    erstellt_am     DATE NOT NULL DEFAULT CURRENT_DATE
);""",
         testfall=[("Testfall: ", {"bold": True}),
                    ("Beide Exemplare des Werks „1984“ -- Eigentümer Hannah Vogel (neu) und Anna Bauer (gut). Zeigt den Mehrwert der Trennung Buch/Buchexemplar aus dem Tutor-Feedback.", {})]),
    dict(title="Tabelle: ZEITSLOT", page=12, shot="zeitslot.png",
         caption="Zwei Zeitslots am 2026-10-02",
         sql="""CREATE TABLE zeitslot (
    zeitslot_id  INTEGER PRIMARY KEY AUTOINCREMENT,
    datum        DATE NOT NULL,
    startzeit    TIME NOT NULL,
    endzeit      TIME NOT NULL,
    CHECK (startzeit < endzeit)
);""",
         testfall=[("Testfall: ", {"bold": True}),
                    ("Freie Zeitslots am 2026-10-02 → 10:00-11:00 und 16:00-17:00.", {})]),
    dict(title="Tabelle: STANDORT", page=13, shot="standort.png",
         caption="Ergebnis: Stadtbibliothek Mitte",
         sql="""CREATE TABLE standort (
    standort_id   INTEGER PRIMARY KEY AUTOINCREMENT,
    bezeichnung   VARCHAR(100) NOT NULL,
    strasse       VARCHAR(100),
    plz           VARCHAR(10),
    stadt         VARCHAR(50) NOT NULL,
    breitengrad   DECIMAL(9,6) NOT NULL CHECK (breitengrad BETWEEN -90 AND 90),
    laengengrad   DECIMAL(9,6) NOT NULL CHECK (laengengrad BETWEEN -180 AND 180)
);""",
         testfall=[("Testfall: ", {"bold": True}),
                    ("Standorte in Berlin → „Stadtbibliothek Mitte“, Breite Strasse 30.", {})]),
    dict(title="Tabelle: VERSANDOPTION", page=14, shot="versandoption.png",
         caption="7 Exemplare mit Versandoption Post/Beides",
         sql="""CREATE TABLE versandoption (
    versandoption_id  INTEGER PRIMARY KEY AUTOINCREMENT,
    exemplar_id       INTEGER NOT NULL UNIQUE
        REFERENCES buchexemplar(exemplar_id) ON DELETE CASCADE,
    versandart        VARCHAR(10) NOT NULL
        CHECK (versandart IN ('abholung','post','beides')),
    kosten            DECIMAL(6,2) NOT NULL DEFAULT 0.00 CHECK (kosten >= 0)
);""",
         testfall=[("Testfall: ", {"bold": True}),
                    ("Alle per Post versendbaren Exemplare (versandart IN ('post','beides')) -- 7 Treffer.", {})]),
    dict(title="Tabelle: BENACHRICHTIGUNG", page=15, shot="benachrichtigung.png",
         caption="Ungelesene Benachrichtigung von Ivo Neumann",
         sql="""CREATE TABLE benachrichtigung (
    benachrichtigung_id  INTEGER PRIMARY KEY AUTOINCREMENT,
    nutzer_id            INTEGER NOT NULL REFERENCES nutzer(nutzer_id) ON DELETE CASCADE,
    typ                   VARCHAR(30) NOT NULL,
    inhalt                TEXT NOT NULL,
    gelesen               BOOLEAN NOT NULL DEFAULT 0 CHECK (gelesen IN (0,1)),
    erstellt_am           DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP
);""",
         testfall=[("Testfall: ", {"bold": True}),
                    ("Ungelesene Benachrichtigungen von Ivo Neumann → „Ein gewuenschtes Buch ist jetzt verfuegbar.“", {})]),
    dict(title="Tabelle: BEWERTUNG", page=16, shot="bewertung.png",
         caption="Ranking der Eigentuemer:innen nach Durchschnittsbewertung",
         sql="""CREATE TABLE bewertung (
    bewertung_id      INTEGER PRIMARY KEY AUTOINCREMENT,
    ausleihe_id       INTEGER NOT NULL UNIQUE
        REFERENCES ausleihe(ausleihe_id) ON DELETE CASCADE,
    sternebewertung   INTEGER NOT NULL CHECK (sternebewertung BETWEEN 1 AND 5),
    kommentar         TEXT,
    bewertungsdatum   DATE NOT NULL DEFAULT CURRENT_DATE
);""",
         testfall=[("Testfall: ", {"bold": True}),
                    ("Durchschnittsbewertung je Eigentümer:in (GROUP BY) -- Top: Greta Klein, Clara Hoffmann, Ben Fischer je 5.0. ausleihe_id ist UNIQUE → höchstens eine Bewertung je Ausleihe (Feedback Tutor Punkt 1).", {})]),
]

for e in entities:
    add_entity_slide(e["title"], e["page"], e["sql"], e["testfall"], e["shot"], e["caption"])

# ---------------------------------------------------------------------------
# Dreifachbeziehungen
# ---------------------------------------------------------------------------
add_entity_slide(
    "Dreifachbeziehung: AUSLEIHE (LEIHT_AUS)", 17, """CREATE TABLE ausleihe (
    ausleihe_id    INTEGER PRIMARY KEY AUTOINCREMENT,
    nutzer_id      INTEGER NOT NULL REFERENCES nutzer(nutzer_id),
    exemplar_id    INTEGER NOT NULL REFERENCES buchexemplar(exemplar_id),
    zeitslot_id    INTEGER NOT NULL REFERENCES zeitslot(zeitslot_id),
    ausleihdatum   DATE NOT NULL,
    rueckgabedatum DATE,
    status VARCHAR(15) NOT NULL DEFAULT 'aktiv'
        CHECK (status IN ('aktiv','zurueckgegeben','ueberfaellig')),
    CHECK (rueckgabedatum IS NULL OR rueckgabedatum >= ausleihdatum),
    UNIQUE (exemplar_id, zeitslot_id)
);""",
    [("Testfall: ", {"bold": True}),
     ("Alle abgeschlossenen Ausleihen mit Entleiher:in und Buchtitel -- 10 Treffer (Nutzer:in × Buchexemplar × Zeitslot).", {})],
    "ausleihe.png", "10 abgeschlossene Ausleihvorgaenge")

add_entity_slide(
    "Dreifachbeziehung: VERFUEGBARKEIT (VERFUEGBAR_AN)", 18, """CREATE TABLE verfuegbarkeit (
    verfuegbarkeit_id  INTEGER PRIMARY KEY AUTOINCREMENT,
    exemplar_id  INTEGER NOT NULL REFERENCES buchexemplar(exemplar_id) ON DELETE CASCADE,
    zeitslot_id  INTEGER NOT NULL REFERENCES zeitslot(zeitslot_id) ON DELETE CASCADE,
    standort_id  INTEGER NOT NULL REFERENCES standort(standort_id),
    verfuegbar   BOOLEAN NOT NULL DEFAULT 1 CHECK (verfuegbar IN (0,1)),
    UNIQUE (exemplar_id, zeitslot_id, standort_id)
);""",
    [("Testfall: ", {"bold": True}),
     ("Verfügbare Exemplare an der „Stadtbibliothek Mitte“ -- kombiniert Buchexemplar, Zeitslot und Standort zu einer Verfügbarkeitsaussage.", {})],
    "verfuegbarkeit.png", "Verfuegbare Exemplare an der Stadtbibliothek Mitte")

# ---------------------------------------------------------------------------
# M:N-Verknuepfungstabellen (kein Screenshot noetig)
# ---------------------------------------------------------------------------
s = new_slide()
add_header(s, "M:N-Verknüpfungstabellen: BUCH_AUTOR und BUCH_GENRE", 19, TOTAL)
add_text(s, Inches(0.5), Inches(1.2), Inches(11), Inches(0.35), [("CREATE-Statements:", {"bold": True, "size": 15})])
add_code_box(s, Inches(0.5), Inches(1.6), Inches(8.5), Inches(3.4), """CREATE TABLE buch_autor (
    buch_id   INTEGER NOT NULL REFERENCES buch(buch_id)   ON DELETE CASCADE,
    autor_id  INTEGER NOT NULL REFERENCES autor(autor_id) ON DELETE CASCADE,
    PRIMARY KEY (buch_id, autor_id)
);
CREATE TABLE buch_genre (
    buch_id   INTEGER NOT NULL REFERENCES buch(buch_id)   ON DELETE CASCADE,
    genre_id  INTEGER NOT NULL REFERENCES genre(genre_id) ON DELETE CASCADE,
    PRIMARY KEY (buch_id, genre_id)
);""", font_size=13)
add_text(s, Inches(0.5), Inches(5.2), Inches(11.5), Inches(1.5),
         "Lösen die binären M:N-Beziehungen Verfasst und Gehört zu aus dem ER-Modell (Phase 1) "
         "relational auf. Zusammengesetzter Primärschlüssel verhindert doppelte Zuordnungen.",
         font_size=15)

# ---------------------------------------------------------------------------
# Trigger
# ---------------------------------------------------------------------------
s = new_slide()
add_header(s, "Trigger für Geschäftsregeln (nicht per CHECK abbildbar)", 20, TOTAL)
add_text(s, Inches(0.5), Inches(1.15), Inches(12.3), Inches(0.7),
         "SQLite erlaubt keine Subqueries in CHECK-Constraints -- Regeln, die auf andere "
         "Tabellen verweisen, werden daher per TRIGGER durchgesetzt:", font_size=15)
add_code_box(s, Inches(0.5), Inches(1.9), Inches(12.3), Inches(4.6), """CREATE TRIGGER trg_ausleihe_kein_eigenausleih
BEFORE INSERT ON ausleihe FOR EACH ROW
WHEN NEW.nutzer_id = (SELECT eigentuemer_id FROM buchexemplar
                       WHERE exemplar_id = NEW.exemplar_id)
BEGIN SELECT RAISE(ABORT, 'Nutzer kann eigenes Buchexemplar nicht ausleihen'); END;

CREATE TRIGGER trg_bewertung_nur_nach_rueckgabe
BEFORE INSERT ON bewertung FOR EACH ROW
WHEN (SELECT status FROM ausleihe WHERE ausleihe_id = NEW.ausleihe_id) != 'zurueckgegeben'
BEGIN SELECT RAISE(ABORT, 'Bewertung nur fuer abgeschlossene Ausleihen erlaubt'); END;""", font_size=14)

# ---------------------------------------------------------------------------
# Negativtestfaelle (2 Folien, 3 Kacheln pro Folie)
# ---------------------------------------------------------------------------
neg_cases_1 = [
    ("(1) Sternebewertung = 7", "negativtest1.png", "CHECK: sternebewertung BETWEEN 1 AND 5"),
    ("(2) Versandkosten = -5.00", "negativtest2.png", "CHECK: kosten >= 0"),
    ("(3) Zeitslot 15:00-14:00", "negativtest3.png", "CHECK: startzeit < endzeit"),
]
neg_cases_2 = [
    ("(4) Ausleihe mit nutzer_id = 9999", "negativtest4.png", "FOREIGN KEY constraint failed (mit PRAGMA foreign_keys = ON;)"),
    ("(5) Eigentümer:in leiht eigenes Exemplar", "negativtest5.png", "Trigger: Nutzer kann eigenes Buchexemplar nicht ausleihen"),
    ("(6) Bewertung zu aktiver Ausleihe", "negativtest6.png", "Trigger: Bewertung nur für abgeschlossene Ausleihen erlaubt"),
]


def add_negtest_slide(title, page_no, cases, note=None):
    s = new_slide()
    add_header(s, title, page_no, TOTAL)
    tile_w = Inches(3.95)
    gap = Inches(0.2)
    top = Inches(1.3) if not note else Inches(1.3)
    for i, (label, fname, expected) in enumerate(cases):
        x = Inches(0.4) + i * (tile_w + gap)
        add_text(s, x, top, tile_w, Inches(0.5), [(label, {"bold": True, "size": 12})], align=PP_ALIGN.CENTER)
        add_screenshot(s, os.path.join(SHOTS, fname), x, top + Inches(0.5), tile_w, Inches(3.0))
        add_text(s, x, top + Inches(3.6), tile_w, Inches(0.9), [(expected, {"size": 11})], align=PP_ALIGN.CENTER)
    if note:
        add_text(s, Inches(0.4), Inches(6.15), Inches(12.5), Inches(1.1),
                 [("Hinweis: ", {"bold": True, "italic": True}), (note, {"italic": True, "size": 13})], font_size=13)
    return s


add_negtest_slide("Negativtestfälle: Geschäftsregeln aktiv geprüft (1/2)", 21, neg_cases_1)
add_negtest_slide("Negativtestfälle: Geschäftsregeln aktiv geprüft (2/2)", 22, neg_cases_2,
                   note="SQLite erzwingt Fremdschlüssel nur, wenn PRAGMA foreign_keys = ON; pro Verbindung gesetzt ist "
                        "-- ohne dieses PRAGMA würde Test (4) fälschlich durchgehen, obwohl das Schema korrekt ist.")

# ---------------------------------------------------------------------------
# Zusammenfassung / Reflexion / Ende
# ---------------------------------------------------------------------------
s = new_slide()
add_header(s, "Zusammenfassung der Implementierung", 23, TOTAL)
add_bullets(s, Inches(0.6), Inches(1.3), Inches(12.1), Inches(5.5), [
    "Das ER-Modell aus Phase 1 (13 Entitäten, 2 Dreifachbeziehungen) wurde vollständig in SQLite als "
    "17 physische Tabellen umgesetzt: 13 Entitätstabellen, 2 Dreifachbeziehungstabellen (ausleihe, "
    "verfuegbarkeit) und 2 M:N-Verknüpfungstabellen (buch_autor, buch_genre).",
    "Alle 17 Tabellen enthalten mindestens 10 realistische Dummy-Datensätze (insgesamt > 190 Zeilen), "
    "erzeugt und verifiziert per Python-sqlite3-Skript.",
    "Datenintegrität wird durch NOT NULL, UNIQUE, Fremdschlüssel (inkl. ON DELETE CASCADE wo fachlich "
    "sinnvoll) sowie CHECK-Constraints für alle relevanten Wertebereiche sichergestellt; zwei Trigger "
    "decken Geschäftsregeln ab, die CHECK allein nicht abbilden kann.",
    "Für jede Entität liegt mindestens ein erfolgreicher Testfall vor; zusätzlich wurden 6 gezielt "
    "ungültige Fälle formuliert und ausgeführt, die alle korrekt abgelehnt wurden -- direkte Umsetzung "
    "des Tutor-Feedbacks aus Phase 1.",
], font_size=16)

s = new_slide()
add_header(s, "Herausforderungen und Erkenntnisse", 24, TOTAL)
add_bullets(s, Inches(0.6), Inches(1.3), Inches(12.1), Inches(5.5), [
    "Werk vs. Exemplar: Erst durch das Tutor-Feedback wurde klar, dass eigentuemer_id nicht auf BUCH "
    "gehört, sondern auf ein eigenes BUCHEXEMPLAR -- sonst könnte ein Buchtitel nur einer Person gehören.",
    "Bewertung ohne Ausleihe verhindern: Eine reine Dreifachbeziehung Bewertender-Buch-Eigentümer hätte "
    "nicht ausgeschlossen, dass ohne echten Ausleihvorgang bewertet wird. Die Kopplung über ausleihe_id "
    "(UNIQUE) löst das strukturell, nicht nur per Anwendungslogik.",
    "SQLite-Grenzen: Keine nativen ENUM-Typen und keine Subqueries in CHECK-Constraints -- das erzwang "
    "saubere CHECK-Listen bzw. Trigger, was die Constraints expliziter macht als ein einfacher ENUM-Typ.",
    "Nächster Schritt (Phase 3): Weitere Indizes für Abfrageperformance sowie Ausbau der Umkreissuche "
    "(Haversine-Berechnung auf Basis von standort).",
], font_size=16)

s = new_slide()
add_text(s, Inches(1.0), Inches(3.1), Inches(11.3), Inches(0.7),
         [("Vielen Dank für die Aufmerksamkeit!", {})], font_size=28, align=PP_ALIGN.CENTER)
add_text(s, Inches(1.0), Inches(3.9), Inches(11.3), Inches(0.5),
         [("Fragen?", {})], font_size=18, align=PP_ALIGN.CENTER)

prs.save(OUT)
print("PPTX gespeichert:", OUT)
print("Anzahl Folien:", len(prs.slides.__iter__.__self__._sldIdLst))

